"""Diapositivas para entender el experimento 1 del examen parcial.

Uso:
    python EP/slides/construir_exp1.py

Sale Experimento1_EP.pptx y .pdf. Las cifras se leen de EP/results/exp1/.
"""
import json
import sys
from pathlib import Path

import pandas as pd
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

aqui = Path(__file__).resolve().parent
ep = aqui.parent
res = ep / "results" / "exp1"
sys.path.insert(0, str(ep.parent / "pc2" / "slides"))
from comun import exportar, nueva_presentacion  # noqa: E402

NEGRO = RGBColor(0, 0, 0)
CONDS = ["A0", "A", "B", "C"]

resumen = pd.read_csv(res / "resumen.csv", index_col=0)
lat = pd.read_csv(res / "latencia.csv", index_col=0)
fallos = pd.read_csv(res / "fallos.csv", index_col=0)
boot = pd.read_csv(res / "bootstrap.csv")
rankings = {}
for l in open(res / "rankings.jsonl", encoding="utf-8"):
    r = json.loads(l)
    rankings[(r["condicion"], r["query_id"])] = r["pasajes"]


def num(x, d=2):
    return f"{round(float(x) + 1e-9, d):.{d}f}".replace(".", ",")


def signo(x, d=2):
    return ("+" if x > 0 else "") + num(x, d)


def pos(cond, q, pasaje):
    p = rankings[(cond, q)]
    return str(p.index(pasaje) + 1) if pasaje in p else "fuera"


prs = nueva_presentacion()


def slide_titulo(titulo):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.3), Inches(12.1), Inches(0.9)
    t.text_frame.text = titulo
    t.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    t.text_frame.paragraphs[0].alignment = 1
    return s


def bullets(s, items, x=0.6, y=1.4, w=12.1, h=5.6, size=20):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        nivel = 1 if it.startswith("  ") else 0
        p.text = ("- " if nivel else "• ") + it.strip()
        p.level = nivel
        p.space_after = Pt(8)
        for r in p.runs:
            r.font.size = Pt(size - 3 * nivel)
    return tb


def tabla(s, datos, x, y, anchos, alto=0.4, size=15, negrita_max=None):
    gf = s.shapes.add_table(len(datos), len(datos[0]), Inches(x), Inches(y), Inches(sum(anchos)),
                            Inches(alto * len(datos)))
    t = gf.table
    for j, a in enumerate(anchos):
        t.columns[j].width = Inches(a)
    for i, fila in enumerate(datos):
        t.rows[i].height = Inches(alto)
        for j, v in enumerate(fila):
            c = t.cell(i, j)
            c.text = str(v)
            c.margin_top = c.margin_bottom = Inches(0.03)
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size)
                    if negrita_max and i > 0 and (i, j) in negrita_max:
                        r.font.bold = True
    return t


def caja(s, x, y, w, h, texto, size=16):
    b = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.background()
    b.line.color.rgb = NEGRO
    b.line.width = Pt(1.25)
    b.shadow.inherit = False
    tf = b.text_frame
    tf.word_wrap = True
    for i, linea in enumerate(texto.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = linea
        p.alignment = 2
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.color.rgb = NEGRO
    return b


def flecha(s, x1, y1, x2, y2):
    from lxml import etree
    c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = NEGRO
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    e = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    e.set("type", "triangle")


def notas(s, texto):
    s.notes_slide.notes_text_frame.text = texto


def negrita(tb, indices):
    for i in indices:
        for r in tb.text_frame.paragraphs[i].runs:
            r.font.bold = True


# 1. portada
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Experimento 1: recuperación y reranking"
s.placeholders[1].text = "Examen Parcial CC0F4 - Tópicos de Ciencia de la Computación I"
s.shapes.title.left, s.shapes.title.top, s.shapes.title.width, s.shapes.title.height = Inches(0.9), Inches(2.0), Inches(11.5), Inches(1.8)
s.placeholders[1].left, s.placeholders[1].top, s.placeholders[1].width, s.placeholders[1].height = Inches(0.9), Inches(4.1), Inches(11.5), Inches(1.0)

# 2. de que trata
s = slide_titulo("¿De qué trata el experimento?")
bullets(s, [
    "Un auditor pregunta por un proceso del SEACE y el sistema debe encontrar los pasajes donde está la respuesta",
    "  Ejemplo: ¿En qué proceso se compró el tóner Xerox con código de referencia 006R01830?",
    "Se comparan tres formas de buscar sobre el mismo material",
    "  A: búsqueda densa (embeddings + FAISS)",
    "  B: A + BM25 + fusión RRF",
    "  C: B + reranking con un Cross-Encoder",
    "Pregunta: ¿cómo cambian la recuperación de evidencia y su costo al agregar cada pieza?",
    "No se busca probar que un método es mejor en general, solo medir qué cambia y cuánto cuesta",
])
notas(s, "Este experimento solo busca. La parte de generar la respuesta con un LLM es el experimento 2, que recibe los "
         "5 primeros resultados de la condición C.")

# 3. que entra
s = slide_titulo("Qué entra al experimento")
bullets(s, [
    "Corpus: 55 procesos reales del SEACE partidos en 426 pasajes",
    "  Un pasaje es una parte del proceso: convocatoria, ítems, cronograma, postores, resultado, contrato",
    "  Cada pasaje tiene un id, por ejemplo E49-P05 (proceso E49, pasaje 5)",
    "22 consultas: 20 con qrels y 2 de control sin evidencia en el corpus (Q21 y Q22)",
    "  Los 2 controles no entran en las métricas de este experimento, los usa el experimento 2",
    "11 de las 20 consultas son difíciles",
    "  Distractores reales (cuatro compras de Office, dos de Adobe), códigos exactos, paráfrasis,",
    "  una consulta numérica, una contradicción y dos que necesitan varios pasajes",
    "Los qrels (qué pasajes son relevantes) se anotaron leyendo los pasajes y se congelaron con SHA-256",
    "  Se fijaron antes de correr cualquier sistema",
], size=19)

# 4. unidad
s = slide_titulo("Unidad de recuperación: el pasaje con su encabezado")
bullets(s, [
    "En A, B y C lo que se busca es un pasaje, con el encabezado de su proceso delante",
    "  Encabezado: nomenclatura y objeto del proceso | entidad",
    "  Así cada pasaje lleva la evidencia y también lo que lo identifica",
    "Referencia A0: la configuración de la PC2",
    "  Chunks de 180 palabras con 1 pasaje de solapamiento (86 chunks en total), solo búsqueda densa",
    "  En la PC2 el chunk podía separar la buena pro del ítem que la identifica (error N5)",
    "A0 se corre para medir el cambio de unidad en vez de suponerlo",
    "Los resultados de los chunks de A0 se pasan a pasajes para compararlos con los mismos qrels",
])

# 5. que se compara
s = slide_titulo("Qué se compara")
caja(s, 4.9, 1.4, 3.5, 0.8, "Consulta del auditor")
flecha(s, 5.5, 2.2, 2.5, 3.0)
flecha(s, 6.65, 2.2, 6.65, 3.0)
flecha(s, 7.8, 2.2, 10.8, 3.0)
caja(s, 0.7, 3.0, 3.6, 1.3, "A\nBúsqueda densa\ne5-small + FAISS")
caja(s, 4.85, 3.0, 3.6, 1.3, "B\nA + BM25 + fusión RRF\n(señal léxica)")
caja(s, 9.0, 3.0, 3.6, 1.3, "C\nB + Cross-Encoder\n(reordena los 20 de B)")
flecha(s, 4.3, 3.65, 4.85, 3.65)
flecha(s, 8.45, 3.65, 9.0, 3.65)
caja(s, 3.4, 5.0, 6.5, 0.9, "Salida de cada condición:\nlista de pasajes ordenados (top-20)")
flecha(s, 2.5, 4.3, 4.5, 5.0)
flecha(s, 6.65, 4.3, 6.65, 5.0)
flecha(s, 10.8, 4.3, 8.8, 5.0)
bullets(s, ["A0 (referencia): la configuración de la PC2, chunks de 180 palabras y solo búsqueda densa"],
        x=0.7, y=6.2, h=0.8, size=17)
notas(s, "Las tres condiciones reciben las mismas consultas y buscan sobre los mismos pasajes. "
         "Cada condición agrega una pieza a la anterior.")

# 6. A
s = slide_titulo("A: búsqueda densa")
bullets(s, [
    "El modelo e5-small convierte la consulta y cada pasaje en un vector de 384 números",
    "FAISS compara el vector de la consulta con los 426 vectores y devuelve los más parecidos (top-20)",
    "Compara significado, no palabras: entiende paráfrasis",
    "  Ejemplo: \"de su propiedad\" y \"equipo propio\" quedan cerca",
    "Su punto débil son los códigos y nombres casi iguales",
    "  Q06: el proceso tiene cinco tóneres Xerox con códigos parecidos y para el embedding todos se parecen",
    "  El pasaje correcto queda en el puesto 4 (no entra al top-3)",
], size=20)
notas(s, "e5 usa los prefijos query: y passage:, texto en minúsculas y vectores normalizados, igual que en la PC2. "
         "IndexFlatIP es búsqueda exacta; se comprobó que da el mismo orden que el producto interno en numpy.")

# 7. B
s = slide_titulo("B: A + BM25 + fusión RRF")
bullets(s, [
    "BM25 busca por palabras exactas y da más peso a las palabras raras",
    "  El código 006r01830 aparece en un solo pasaje, así que BM25 lo detecta enseguida",
    "RRF junta la lista del denso y la de BM25 en una sola, usando la posición de cada pasaje",
    "  puntaje = suma de 1 / (60 + posición) en cada lista",
    "  Ejemplo: puesto 4 en el denso y puesto 1 en BM25 da 1/64 + 1/61 = 0,032",
    "  Un pasaje que solo aparece en una lista, en el puesto 1, da 1/61 = 0,016",
    "Así sube lo que ambas listas consideran bueno, sin comparar puntajes de escalas distintas",
    "De A a B cambian dos cosas a la vez (la señal léxica y la fusión)",
    "  No se puede saber cuál de las dos produce cada cambio",
], size=19)

# 8. C
s = slide_titulo("C: B + Cross-Encoder")
bullets(s, [
    "El Cross-Encoder lee la consulta y el pasaje juntos y decide qué tan bien se corresponden",
    "  El denso, en cambio, vectoriza la consulta y el pasaje por separado",
    "Solo reordena los 20 candidatos que entrega B",
    "  No puede recuperar un pasaje que B no trajo",
    "  Por eso B frente a C mide el efecto incremental del reranker",
    "Es más fino pero más lento: hay que correrlo una vez por cada candidato (20 pares por consulta)",
    "Modelo: mmarco-mMiniLMv2, multilingüe, entrenado con datos de búsqueda",
], size=20)
notas(s, "El par más largo que ve el Cross-Encoder tiene 285 tokens, dentro de su ventana de 512.")

# 9. fijo
s = slide_titulo("Qué se mantiene igual")
bullets(s, [
    "Mismos 426 pasajes y mismas 20 consultas con qrels (benchmark congelado)",
    "Mismo modelo de embeddings (e5-small con prefijos query:/passage:, minúsculas y normalización L2)",
    "Misma profundidad: cada recuperador entrega 20 candidatos",
    "BM25 con k1 = 1,5 y b = 0,75; RRF con k = 60",
    "Misma semilla (42), misma máquina (CPU, 8 hilos) y una llamada de calentamiento antes de medir",
    "Por eso las diferencias entre A, B y C se pueden atribuir a la pieza que se agregó",
])

# 10. como se mide
s = slide_titulo("Cómo se mide")
bullets(s, [
    "Recall@k: de los pasajes relevantes, qué fracción aparece en los k primeros (k = 1, 3, 5)",
    "MRR: 1 dividido entre la posición del primer pasaje relevante (puesto 1 = 1; puesto 4 = 0,25)",
    "nDCG@5: premia que los relevantes estén arriba, no solo que estén",
    "Latencia por etapa (media y p95): denso, BM25, RRF y Cross-Encoder",
    "Bootstrap pareado (5000 remuestreos): intervalo de confianza de la diferencia entre dos condiciones",
    "Tipo de fallo por consulta: acierta, recupera parte, mal ordenado o no lo trae",
    "Los qrels son binarios (relevante o no) y están sobre el id del pasaje",
], size=19)

# 11. ejemplo de metricas
s = slide_titulo("Ejemplo de métricas: consulta Q06")
rel = "E49-P05"
datos = [["Condición", "Posición del pasaje E49-P05", "Recall@3", "MRR"]]
for c in CONDS:
    fila = pd.read_csv(res / "por_consulta.csv")
    f = fila[(fila.condicion == c) & (fila.query_id == "Q06")].iloc[0]
    datos.append([c, pos(c, "Q06", rel), num(f["recall@3"]), num(f["mrr"])])
tabla(s, datos, 0.9, 1.6, [2.4, 4.6, 2.3, 2.0], alto=0.5, size=18)
bullets(s, [
    "Q06 tiene un solo pasaje relevante",
    "  Si queda en el puesto 4 o 5, Recall@3 vale 0, pero Recall@5 sí lo cuenta",
    "  MRR baja a 0,25 o 0,2: sigue premiando que esté cerca de arriba",
    "En B y C el pasaje sube al puesto 1 y las dos métricas valen 1",
    "Cuando una consulta tiene dos pasajes relevantes, Recall@1 no puede pasar de 0,5",
    "  Por eso el Recall@1 promedio de C (0,61) es menor que su MRR (0,88)",
], y=4.5, h=2.6, size=18)

# 12. resultados: tabla
s = slide_titulo("Resultados: métricas")
nombres = {"A0": "A0 (PC2)", "A": "A (denso)", "B": "B (A + BM25 + RRF)", "C": "C (B + Cross-Encoder)"}
cols = ["recall@1", "recall@3", "recall@5", "mrr", "ndcg@5"]
datos = [["Condición", "Recall@1", "Recall@3", "Recall@5", "MRR", "nDCG@5"]]
for c in CONDS:
    datos.append([nombres[c]] + [num(resumen.loc[c, m]) for m in cols])
mejores = {(CONDS.index(resumen[m].idxmax()) + 1, j + 1) for j, m in enumerate(cols)}
tabla(s, datos, 0.7, 1.6, [3.9, 1.6, 1.6, 1.6, 1.6, 1.6], alto=0.55, size=18, negrita_max=mejores)
bullets(s, [
    "En negrita, el mejor valor de cada columna",
    "Media sobre las 20 consultas con qrels",
    "C gana en Recall@1, Recall@3, MRR y nDCG@5; B gana en Recall@5",
], y=4.5, h=2.2, size=20)

# 13. resultados: grafico
s = slide_titulo("Resultados: gráfico")
cd = CategoryChartData()
cd.categories = ["Recall@1", "Recall@3", "Recall@5", "MRR", "nDCG@5"]
for c in CONDS:
    cd.add_series(c, [float(resumen.loc[c, m]) for m in cols])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.4), Inches(8.0), Inches(5.4), cd)
ch = gf.chart
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
pl = ch.plots[0]
pl.gap_width = 60
for punto, m in zip(pl.series[3].points, cols):
    tf = punto.data_label.text_frame
    tf.text = num(resumen.loc["C", m])
    tf.paragraphs[0].runs[0].font.size = Pt(11)
ch.value_axis.maximum_scale = 1.0
ch.value_axis.minimum_scale = 0
ch.value_axis.has_major_gridlines = False
bullets(s, [
    "Cada grupo de barras es una métrica",
    "Cada color es una condición",
    "Se rotulan los valores de C",
    "A0 es la más baja en todo",
], x=8.9, y=1.8, w=4.0, h=4.5, size=18)

# 14. interpretacion
s = slide_titulo("Interpretación de las métricas")
bullets(s, [
    "La unidad importó: de A0 a A el MRR sube de " + num(resumen.loc["A0", "mrr"]) + " a " + num(resumen.loc["A", "mrr"]),
    "  Solo se cambió de chunks de 180 palabras a pasaje con encabezado, con el mismo modelo",
    "B casi no se distingue de A",
    "  Recall@5 sube de " + num(resumen.loc["A", "recall@5"]) + " a " + num(resumen.loc["B", "recall@5"]) + " y el MRR apenas de "
    + num(resumen.loc["A", "mrr"]) + " a " + num(resumen.loc["B", "mrr"]),
    "C mejora sobre todo el principio de la lista",
    "  Recall@1 sube de " + num(resumen.loc["B", "recall@1"]) + " a " + num(resumen.loc["C", "recall@1"])
    + " y el MRR de " + num(resumen.loc["B", "mrr"]) + " a " + num(resumen.loc["C", "mrr"]),
    "Recall@5 de C (" + num(resumen.loc["C", "recall@5"]) + ") es menor que el de B (" + num(resumen.loc["B", "recall@5"]) + ")",
    "  El reranker ordena mejor, pero en algunas consultas saca del top-5 un pasaje que B tenía dentro",
], size=20)

# 15. latencia
s = slide_titulo("Costo: latencia por etapa")
datos = [["Etapa", "Tiempo medio (ms)"],
         ["Denso (e5 + FAISS)", num(lat.loc["B", "dense_ms"], 1)],
         ["BM25", num(lat.loc["B", "bm25_ms"], 1)],
         ["RRF", num(lat.loc["B", "rrf_ms"], 2)],
         ["Cross-Encoder (20 pares)", num(lat.loc["C", "reranker_ms"], 0)]]
tabla(s, datos, 0.7, 1.6, [4.2, 2.6], alto=0.5, size=17)
datos = [["Condición", "Total medio (ms)", "p95 (ms)"]]
for c in CONDS:
    datos.append([c, num(lat.loc[c, "total_ms"], 0), num(lat.loc[c, "total_p95_ms"], 0)])
tabla(s, datos, 8.0, 1.6, [1.6, 2.0, 1.4], alto=0.5, size=17)
razon = lat.loc["C", "total_ms"] / lat.loc["B", "total_ms"]
bullets(s, [
    "BM25 y RRF casi no cuestan: B tarda lo mismo que A",
    f"El Cross-Encoder domina el costo: C tarda unas {num(razon, 0)} veces lo que B",
    "  Hay que evaluar 20 pares consulta-pasaje por cada consulta",
    "Medido en CPU (AMD Ryzen 7, 8 hilos) y una sola corrida, así que sirve de orden de magnitud",
], y=4.6, h=2.5, size=19)

# 16. bootstrap
s = slide_titulo("¿Las diferencias son reales? Bootstrap")


def celda(a, b, m):
    f = boot[(boot.comparacion == f"{b} - {a}") & (boot.metrica == m)].iloc[0]
    return f"{signo(f.delta_media, 3)} [{num(f.ic95_bajo, 3)}; {num(f.ic95_alto, 3)}]"


datos = [["Comparación", "Recall@3", "MRR", "nDCG@5"]]
for a, b in [("A", "B"), ("B", "C"), ("A", "C"), ("A0", "A")]:
    datos.append([f"{b} - {a}"] + [celda(a, b, m) for m in ("recall@3", "mrr", "ndcg@5")])
tabla(s, datos, 0.7, 1.5, [2.2, 3.3, 3.3, 3.3], alto=0.5, size=16)
bullets(s, [
    "Cada celda: diferencia de medias y su intervalo de confianza al 95 %",
    "Si el intervalo incluye el 0, con 20 consultas no se puede distinguir la diferencia de cero",
    "B - A y C - B cruzan el cero en las tres métricas",
    "C - A es la única comparación con las tres métricas por encima de cero",
    "A - A0 también, pero solo en MRR y nDCG@5",
], y=4.3, h=2.8, size=19)
notas(s, "Que el promedio de C sea más alto que el de B no basta para decir que C es mejor: el intervalo de C - B "
         "todavía incluye el cero.")

# 17. fallos
s = slide_titulo("Tipos de fallo por consulta")
cab = ["RETRIEVAL_OK", "PARTIAL_RECALL", "RANKING_ERROR", "RETRIEVAL_MISS"]
datos = [["Condición"] + cab]
for c in CONDS:
    datos.append([c] + [str(int(fallos.loc[c, k])) if k in fallos.columns else "0" for k in cab])
tabla(s, datos, 0.7, 1.5, [2.0, 2.5, 2.6, 2.6, 2.4], alto=0.5, size=16)
bullets(s, [
    "Corte en k = 3, sobre las 20 consultas",
    "  RETRIEVAL_OK: todos los pasajes relevantes están en el top-3",
    "  PARTIAL_RECALL: algunos están en el top-3, otros no",
    "  RANKING_ERROR: están entre los candidatos, pero ninguno en el top-3",
    "  RETRIEVAL_MISS: ninguno está entre los candidatos",
    "De A0 a C los aciertos pasan de " + str(int(fallos.loc["A0", "RETRIEVAL_OK"])) + " a " + str(int(fallos.loc["C", "RETRIEVAL_OK"]))
    + " y los errores de orden de " + str(int(fallos.loc["A0", "RANKING_ERROR"])) + " a " + str(int(fallos.loc["C", "RANKING_ERROR"])),
], y=4.3, h=2.8, size=19)

# 18. consultas donde ayudan
s = slide_titulo("Consultas analizadas: donde BM25 y el reranker ayudan")
datos = [["Consulta", "Pasaje relevante", "A", "B", "C"]]
for q, p in [("Q06", "E49-P05"), ("Q20", "E12-P06"), ("Q12", "E50-P05")]:
    datos.append([q, p] + [pos(c, q, p) for c in ("A", "B", "C")])
tabla(s, datos, 0.7, 1.5, [1.8, 3.2, 1.6, 1.6, 1.6], alto=0.5, size=17)
tb = bullets(s, [
    "Posición del pasaje relevante en cada condición (1 = primero)",
    "Q06 (código 006R01830): BM25 pesa el código, que está en un solo pasaje, y B lo sube del puesto 4 al 1",
    "Q20 (antivirus del Programa Juntos): A y B dejan la buena pro en los puestos 4 y 5; C la sube al 1",
    "Q12 (único postor del tomógrafo Siemens): el pasaje de postores pasa del puesto 3 al 1 con C",
    "  El Cross-Encoder ve \"único postor\" junto a \"Número de postores registrados: 1\"",
], y=3.7, h=3.5, size=18)

# 19. consultas donde fallan
s = slide_titulo("Consultas analizadas: donde fallan")
datos = [["Consulta", "Pasaje", "A", "B", "C"]]
for q, p in [("Q03", "E45-P02"), ("Q17", "E37-P06"), ("Q13", "E48-P06")]:
    datos.append([q, p] + [pos(c, q, p) for c in ("A", "B", "C")])
tabla(s, datos, 0.7, 1.5, [1.8, 3.2, 1.6, 1.6, 1.6], alto=0.5, size=17)
bullets(s, [
    "Q03 (Adobe admitiendo equivalente): C baja el ítem del puesto 1 al 9",
    "  Prefiere pasajes de Adobe sin equivalente (el distractor): premia \"Adobe\" y no pesa la condición",
    "Q17 (Office 365 del FONDEPES): B baja el pasaje del puesto 3 al 4 al adelantar a INACAL (también Office 365)",
    "  C lo deja en el puesto 2, todavía detrás del distractor",
    "Q13 (adjudicado más de 30 % sobre el valor referencial): falla en todas",
    "  Hay que comparar números y ningún recuperador lo hace; B y C ni siquiera lo traen entre los 20",
], y=3.7, h=3.5, size=18)

# 20. negativos
s = slide_titulo("Resultados negativos")
bullets(s, [
    "El Recall@5 de C es menor que el de B",
    "  En Q03 y Q14 el reranker empuja fuera del top-5 un pasaje que B tenía dentro",
    "B no se distingue de A",
    "  BM25 gana en Q06 y Q02, pero la fusión baja el pasaje correcto en otras (Q17, Q10, Q11)",
    "  En MRR, B empeora en 6 consultas y mejora en 3",
    "C - B también cruza el cero: solo C - A es distinguible de cero",
    "Q13 no la resuelve ninguna condición",
    "Se conservan porque limitan lo que se puede afirmar del experimento",
])

# 21. conclusion
s = slide_titulo("Conclusión")
bullets(s, [
    "C es la mejor en Recall@1, Recall@3, MRR y nDCG@5, pero la diferencia es clara solo frente a A, no frente a B",
    "  Y cuesta cerca de 1 segundo por consulta en CPU, unas 46 veces lo que B",
    "El cambio más barato y con efecto claro fue la unidad de recuperación (A frente a A0)",
    "  El error N5 de la PC2 era de segmentación, no del modelo",
    "BM25 + RRF ayuda con códigos exactos (Q06) y perjudica cuando adelanta un distractor (Q17)",
    "No se declara que C sea mejor en general",
    "  Son 20 consultas anotadas por nosotros y el intervalo de C - B incluye el cero",
    "Los 5 primeros resultados de C son la evidencia que recibe el LLM en el experimento 2",
], size=20)

salida = aqui / "Experimento1_EP.pptx"
prs.save(salida)
exportar(salida, aqui / "Experimento1_EP.pdf", aqui / "_png_exp1")
print("listo", salida, "| diapositivas:", len(prs.slides))
