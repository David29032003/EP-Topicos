"""Diapositivas para entender el experimento 2 del examen parcial.

Uso:
    python EP/slides/construir_exp2.py

Sale Experimento2_EP.pptx y .pdf. Las cifras se leen de EP/results/exp2/.
"""
import sys
from pathlib import Path

import pandas as pd
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

aqui = Path(__file__).resolve().parent
ep = aqui.parent
res = ep / "results" / "exp2"
sys.path.insert(0, str(ep.parent / "pc2" / "slides"))
from comun import exportar, nueva_presentacion  # noqa: E402

NEGRO = RGBColor(0, 0, 0)

resumen = pd.read_csv(res / "resumen.csv", index_col=0)
cb, rg = resumen.loc["closed_book"], resumen.loc["rag_C"]


def num(x, d=2):
    return f"{round(float(x) + 1e-9, d):.{d}f}".replace(".", ",")


prs = nueva_presentacion()


def slide_titulo(titulo):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.3), Inches(12.1), Inches(0.9)
    t.text_frame.text = titulo
    t.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    t.text_frame.paragraphs[0].alignment = 1
    return s


def bullets(s, items, x=0.6, y=1.4, w=12.1, h=5.6, size=20):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        nivel = 1 if it.startswith("  ") else 0
        p.text = ("- " if nivel else "• ") + it.strip()
        p.level = nivel
        p.space_after = Pt(8)
        for r in p.runs:
            r.font.size = Pt(size - 3 * nivel)
    return tb


def tabla(s, datos, x, y, anchos, alto=0.4, size=15):
    gf = s.shapes.add_table(len(datos), len(datos[0]), Inches(x), Inches(y), Inches(sum(anchos)),
                            Inches(alto * len(datos)))
    t = gf.table
    for j, a in enumerate(anchos):
        t.columns[j].width = Inches(a)
    for i, fila in enumerate(datos):
        t.rows[i].height = Inches(alto)
        for j, v in enumerate(fila):
            c = t.cell(i, j)
            c.text = str(v)
            c.margin_top = c.margin_bottom = Inches(0.03)
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size)
    return t


def caja(s, x, y, w, h, texto, size=16):
    b = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.background()
    b.line.color.rgb = NEGRO
    b.line.width = Pt(1.25)
    b.shadow.inherit = False
    tf = b.text_frame
    tf.word_wrap = True
    for i, linea in enumerate(texto.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = linea
        p.alignment = 2
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.color.rgb = NEGRO
    return b


def flecha(s, x1, y1, x2, y2):
    from lxml import etree
    c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = NEGRO
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    e = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    e.set("type", "triangle")


def notas(s, texto):
    s.notes_slide.notes_text_frame.text = texto


# 1. portada
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Experimento 2: RAG y evidencia"
s.placeholders[1].text = "Examen Parcial CC0F4 - Tópicos de Ciencia de la Computación I"
s.shapes.title.left, s.shapes.title.top, s.shapes.title.width, s.shapes.title.height = Inches(0.9), Inches(2.0), Inches(11.5), Inches(1.8)
s.placeholders[1].left, s.placeholders[1].top, s.placeholders[1].width, s.placeholders[1].height = Inches(0.9), Inches(4.1), Inches(11.5), Inches(1.0)

# 2. de que trata
s = slide_titulo("¿De qué trata el experimento?")
bullets(s, [
    "El experimento 1 buscó los pasajes. Aquí un LLM usa esos pasajes para responder",
    "  RAG: el modelo recibe la pregunta y los 5 mejores pasajes de la condición C, y responde citándolos",
    "Se compara contra closed-book: el mismo modelo, con la misma pregunta, pero sin pasajes",
    "Preguntas del experimento",
    "  ¿Entregar evidencia mejora la corrección y el sustento de las respuestas?",
    "  ¿El modelo se abstiene cuando no hay evidencia?",
    "La medida principal es una auditoría manual: leímos las respuestas y los pasajes",
])
notas(s, "RAG significa Retrieval-Augmented Generation: primero se recupera evidencia y luego se genera la respuesta con ella.")

# 3. que entra
s = slide_titulo("Qué entra al experimento")
bullets(s, [
    "12 consultas fijadas antes de correr (marcadas exp2 en consultas.jsonl)",
    "  10 tienen respuesta en el corpus: Q01, Q02, Q04, Q06, Q09, Q10, Q12, Q14, Q15 y Q17",
    "  2 son controles sin evidencia: Q21 (sanciones) y Q22 (plazo de entrega)",
    "  En los controles lo correcto es abstenerse, porque el corpus no tiene ese dato",
    "Modelo: Qwen2.5-1.5B-Instruct, con la revisión fijada",
    "Contexto RAG: los 5 primeros pasajes de la condición C del experimento 1",
    "Respuestas de referencia (referencias.jsonl): se usan para auditar y se congelaron antes de correr",
], size=19)

# 4. que se compara
s = slide_titulo("Qué se compara")
caja(s, 4.9, 1.25, 3.5, 0.7, "Pregunta del auditor", size=15)
flecha(s, 5.6, 1.95, 3.75, 2.5)
flecha(s, 7.7, 1.95, 9.55, 2.5)
caja(s, 1.6, 2.5, 4.3, 1.1, "Closed-book\nContexto vacío\n(no se entregan pasajes)", size=15)
caja(s, 7.4, 2.5, 4.3, 1.1, "RAG\nContexto = top-5 de C\n(5 pasajes con su id)", size=15)
flecha(s, 3.75, 3.6, 5.8, 4.1)
flecha(s, 9.55, 3.6, 7.5, 4.1)
caja(s, 3.9, 4.1, 5.5, 0.9, "Mismo prompt y mismo LLM\nQwen2.5-1.5B, greedy, semilla 42", size=15)
flecha(s, 6.65, 5.0, 6.65, 5.3)
caja(s, 3.9, 5.3, 5.5, 0.7, "Respuesta JSON: respuesta, citas, abstención", size=15)
flecha(s, 6.65, 6.0, 6.65, 6.3)
caja(s, 3.9, 6.3, 5.5, 0.7, "Auditoría manual", size=15)
notas(s, "Lo único que cambia entre las dos condiciones es el bloque de contexto. Por eso cualquier diferencia se puede "
         "atribuir a la evidencia.")

# 5. el prompt
s = slide_titulo("El prompt")
bullets(s, [
    "El mensaje de sistema fija cuatro reglas",
    "  1. Si hay pasajes, usa solo esos y pon en \"citas\" los ids de los que sustentan la respuesta",
    "  2. Si no hay pasajes, responde con lo que sepas y deja \"citas\" vacía",
    "  3. Si no tienes la información, no adivines: \"abstencion\": true y explica qué falta",
    "  4. No inventes identificadores, nombres, montos ni fechas",
    "El mensaje del usuario trae la pregunta y el contexto",
    "  Cada pasaje empieza con su id entre corchetes, por ejemplo [E49-P05], seguido del encabezado del proceso y su texto",
    "El mismo prompt se usa en las dos condiciones; en closed-book el contexto dice \"(no se entregan pasajes)\"",
], size=19)

# 6. la salida
s = slide_titulo("La salida del modelo")
caja(s, 0.9, 1.5, 11.5, 1.9,
     "{\"respuesta\": \"El tóner Xerox 006R01830 fue comprado en el proceso RES-PROC-161-2025-RE-1, ítem 4\",\n"
     "\"citas\": [\"E49-P05\"],\n\"abstencion\": false}", size=16)
bullets(s, [
    "Ejemplo real: respuesta del modelo con RAG a la consulta Q06",
    "respuesta: el texto breve en español",
    "citas: los ids de los pasajes que dice usar",
    "abstencion: true si decide que no tiene la información",
    "El formato fijo permite revisar automáticamente si el JSON es válido y si las citas existen",
    "Pero un JSON válido no dice si la respuesta es verdadera: eso lo revisa la auditoría",
], y=3.7, h=3.4, size=19)

# 7. fijo
s = slide_titulo("Qué se mantiene igual")
bullets(s, [
    "Mismo modelo: Qwen2.5-1.5B-Instruct con la revisión fijada",
    "Decoding greedy (do_sample = False), repetition_penalty = 1,0, máximo 192 tokens nuevos, semilla 42",
    "Mismo prompt, guardado antes de generar las respuestas",
    "Mismas 12 preguntas",
    "El formato del JSON se revisó antes solo con Q05 y Q16, que no forman parte del experimento",
    "  Así no se ajustó el prompt mirando las consultas que se evalúan",
    "Con greedy no hay azar: la diferencia entre closed-book y RAG se explica por el contexto",
], size=19)

# 8. por que el mismo prompt
s = slide_titulo("Por qué el mismo prompt en las dos condiciones")
bullets(s, [
    "Si closed-book tuviera la regla \"usa solo los pasajes\", se abstendría siempre",
    "  La comparación no diría nada, porque no habría respuesta que comparar",
    "Por eso la regla dice: si hay pasajes, usa solo esos y cítalos; si no hay, responde con lo que sepas",
    "  En ambos casos, abstente si no tienes la información",
    "Así las dos condiciones pueden responder o abstenerse, y se compara cómo lo hacen",
    "Por qué top-5 y no top-3",
    "  Varias respuestas necesitan dos pasajes del mismo proceso (postores y buena pro, objeto e ítem)",
    "  C tiene Recall@5 = 0,90 frente a Recall@3 = 0,88, y se fijó antes de correr",
], size=19)

# 9. como se mide
s = slide_titulo("Cómo se mide")
bullets(s, [
    "Corrección de la respuesta (manual, contra la referencia)",
    "  correcta, parcial, incorrecta (afirma algo falso), no responde (se abstuvo aunque la evidencia estaba) o abstención correcta (control)",
    "Corrección de las citas",
    "  Automático: ¿la cita está en el contexto? ¿está en los qrels?",
    "  Manual: ¿la cita sustenta alguna afirmación? ¿la afirmación cita el pasaje que la sustenta?",
    "Faithfulness",
    "  La respuesta se divide en afirmaciones sueltas y cada una se marca como sustentada, no sustentada o contradicha",
    "  Se verifica contra el contexto entregado (en closed-book, contra el corpus)",
    "No se usa fluidez ni la opinión de otro LLM como prueba de calidad",
], size=18)

# 10. ejemplo de auditoria
s = slide_titulo("Ejemplo de auditoría: consulta Q12")
datos = [["Pregunta", "¿Quién fue el único postor en el mantenimiento del tomógrafo Siemens del Hospital General de Jaén?"],
         ["Respuesta del modelo", "SIEMENS HEALTHCARE S.A.C. fue el único postor"],
         ["Cita del modelo", "E50-P01"],
         ["Pasaje que sí lo sustenta", "E50-P05 (postores), que estaba en el puesto 1 del contexto"]]
tabla(s, datos, 0.7, 1.5, [3.4, 8.5], alto=0.65, size=16)
bullets(s, [
    "Corrección de la respuesta: correcta",
    "Corrección de la cita: equivocada",
    "  E50-P01 es la convocatoria por proveedor único; no nombra a la empresa ni cuenta postores",
    "  La respuesta acierta, pero la cita no permite comprobarla",
    "Por eso se auditan por separado la respuesta y sus citas",
], y=4.4, h=2.8, size=19)

# 11. resultados: correccion
s = slide_titulo("Resultados: corrección de las respuestas")
datos = [["", "Closed-book", "RAG (top-5 de C)"],
         ["JSON exacto y con formato", f"{int(cb.formato_valido)}/12", f"{int(rg.formato_valido)}/12"],
         ["Correctas", int(cb.correctas), int(rg.correctas)],
         ["Parciales", int(cb.parciales), int(rg.parciales)],
         ["Incorrectas (afirman algo falso)", int(cb.incorrectas), int(rg.incorrectas)],
         ["No responde (había evidencia)", int(cb.no_responde), int(rg.no_responde)],
         ["Abstención correcta en los 2 controles", cb.abstencion_correcta_en_controles, rg.abstencion_correcta_en_controles]]
tabla(s, datos, 0.7, 1.5, [6.0, 2.8, 3.1], alto=0.52, size=17)
bullets(s, [
    "12 consultas por condición: 10 con respuesta y 2 controles",
    "La auditoría de cada respuesta está en auditoria_respuestas.csv",
], y=5.4, h=1.6, size=19)

# 12. grafico
s = slide_titulo("Resultados: gráfico")
cd = CategoryChartData()
cd.categories = ["Correctas", "Parciales", "Incorrectas", "No responde", "Abstención correcta"]
cd.add_series("Closed-book", [int(cb.correctas), int(cb.parciales), int(cb.incorrectas), int(cb.no_responde), 2])
cd.add_series("RAG", [int(rg.correctas), int(rg.parciales), int(rg.incorrectas), int(rg.no_responde), 0])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.4), Inches(8.2), Inches(5.4), cd)
ch = gf.chart
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
pl = ch.plots[0]
pl.gap_width = 60
pl.has_data_labels = True
pl.data_labels.show_value = True
pl.data_labels.font.size = Pt(14)
ch.category_axis.tick_labels.font.size = Pt(12)
ch.value_axis.maximum_scale = 10
ch.value_axis.minimum_scale = 0
ch.value_axis.has_major_gridlines = False
bullets(s, [
    "Número de consultas de las 12",
    "Closed-book casi nunca responde",
    "RAG responde casi todo, a veces mal",
    "En los controles ocurre lo contrario",
], x=9.1, y=1.8, w=3.8, h=4.5, size=18)

# 13. interpretacion de la correccion
s = slide_titulo("Interpretación: la corrección")
bullets(s, [
    "Sin evidencia el modelo casi no responde: 9 de 10 consultas sin respuesta",
    "  Se abstiene (8 veces) o dice que no tiene información (en Q09, con abstención en false: JSON válido pero incoherente)",
    "  Solo inventa una vez: en Q01 dice que la entidad fue el MTC, y fue la Agencia de Compras de las FF.AA.",
    "Con evidencia responde bien 8 de 10 y a medias 1",
    "  Pasa de no saber nada del SEACE a encontrar la respuesta en los pasajes",
    "Pero RAG pierde la abstención: 0 de 2 controles, frente a 2 de 2 en closed-book",
    "  El recuperador siempre entrega 5 pasajes, aunque ninguno sea relevante, y el modelo los usa para responder",
], size=19)

# 14. citas
s = slide_titulo("Resultados: las citas")
datos = [["Medida", "RAG"],
         ["Citas dentro del contexto entregado (automático)", rg.citas_en_contexto_auto],
         ["Citas que están en los qrels (automático)", rg.citas_en_qrels_auto],
         ["Citas que sustentan alguna afirmación (manual)", rg.citas_que_sustentan_manual],
         ["Afirmaciones sustentadas que citan su pasaje (manual)", rg.afirmaciones_sustentadas_con_su_cita]]
tabla(s, datos, 0.7, 1.5, [8.6, 3.2], alto=0.52, size=17)
bullets(s, [
    "Closed-book no cita nada, porque no tiene pasajes",
    "Todas las citas apuntan a un pasaje que sí estaba en el contexto, pero solo 12 de 25 sustentan algo",
    "  A veces el modelo cita los 5 pasajes del contexto aunque solo 1 o 2 dicen lo que afirma (Q02 y Q14)",
    "Los dos conteos de 12/25 no cuentan las mismas citas",
    "  En Q04 la cita E43-P05 está en los qrels, pero no sustenta el ganador ni el monto",
    "  En Q15 la cita E36-P06 sustenta el 79 %, pero no está en los qrels",
], y=4.4, h=2.9, size=18)

# 15. faithfulness
s = slide_titulo("Resultados: faithfulness")
datos = [["Condición", "Afirmaciones", "Sustentadas", "Faithfulness"],
         ["Closed-book", int(cb.afirmaciones), int(cb.afirmaciones_sustentadas), num(cb.faithfulness)],
         ["RAG (top-5 de C)", int(rg.afirmaciones), int(rg.afirmaciones_sustentadas), num(rg.faithfulness)]]
tabla(s, datos, 0.7, 1.5, [4.3, 2.5, 2.5, 2.6], alto=0.55, size=18)
bullets(s, [
    "Faithfulness = afirmaciones sustentadas dividido entre afirmaciones totales",
    "Con RAG, 16 de 19 afirmaciones están respaldadas por el contexto",
    "Las 3 que no lo están son de las consultas Q15, Q21 y Q22",
    "  Q15: contradicha por el pasaje; Q21 y Q22: ningún pasaje las respalda",
    "Closed-book casi no afirma nada (solo la del MTC en Q01, que es falsa)",
    "  Su faithfulness no se puede comparar, porque casi no hay afirmaciones que revisar",
], y=3.5, h=3.6, size=19)

# 16. errores: controles
s = slide_titulo("Errores analizados: los controles Q21 y Q22")
datos = [["Consulta", "Qué dice el contexto", "Respuesta del modelo con RAG"],
         ["Q21: ¿qué sanción le impuso el Tribunal a MC INNOVATEC?",
          "E01-P06: buena pro a MC INNOVATEC por S/ 184,500.00",
          "Impuso una sanción de S/ 184,500.00 (cita E01-P06)"],
         ["Q22: ¿cuál era el plazo de entrega de los switches Cisco?",
          "Ningún pasaje habla de plazos",
          "Fue de 30 días (cita E01-P02)"]]
tabla(s, datos, 0.6, 1.5, [4.3, 3.8, 4.0], alto=0.95, size=15)
bullets(s, [
    "En ambas debía abstenerse y no lo hizo",
    "Q21: convierte una adjudicación en una sanción, y la cita hace que la respuesta parezca sustentada",
    "Q22: inventa un dato (los plazos están en las bases en PDF, que el corpus no tiene)",
    "Closed-book, sin pasajes, se abstuvo en las dos",
], y=4.7, h=2.4, size=19)

# 17. errores: otros
s = slide_titulo("Errores analizados: Q15, Q12 y Q04")
bullets(s, [
    "Q15: el modelo se contradice",
    "  Dice \"no aceptaba un producto equivalente\" y en la frase siguiente \"lo solicitado era AutoCAD o equivalente\"",
    "  El objeto del proceso dice \"o equivalente\" y el ítem no; el modelo no señala esa diferencia",
    "Q12: respuesta correcta con cita equivocada",
    "  Cita la convocatoria (E50-P01) y no el pasaje de postores (E50-P05), que sí estaba en el contexto",
    "Q04: respuesta parcial",
    "  Da el ganador y el monto, pero omite los 15 postores y cita el pasaje de postores para el ganador",
    "La métrica automática cuenta esa cita de Q04 como correcta porque el pasaje está en los qrels",
], size=19)

# 18. costo
s = slide_titulo("Costo del RAG")
datos = [["", "Closed-book", "RAG"],
         ["Tokens de prompt (media)", int(cb.prompt_tokens_medio), int(rg.prompt_tokens_medio)],
         ["Latencia media por respuesta (s)", num(cb.latencia_media_ms / 1000, 1), num(rg.latencia_media_ms / 1000, 1)],
         ["Latencia p95 (s)", num(cb.latencia_p95_ms / 1000, 1), num(rg.latencia_p95_ms / 1000, 1)]]
tabla(s, datos, 0.7, 1.5, [6.0, 2.9, 2.9], alto=0.55, size=18)
bullets(s, [
    f"Los 5 pasajes multiplican la entrada por {num(rg.prompt_tokens_medio / cb.prompt_tokens_medio, 1)}",
    f"La latencia sube un {num((rg.latencia_media_ms / cb.latencia_media_ms - 1) * 100, 0)} %",
    "  El modelo tiene que leer el contexto antes de generar la primera palabra",
    "Medido en CPU y en una sola corrida: sirve de orden de magnitud",
], y=4.0, h=3.0, size=20)

# 19. lo que no funciono
s = slide_titulo("Lo que no funcionó: abstenerse con un umbral")
datos = [["Consulta", "Puntaje del Cross-Encoder del primer pasaje", "¿Hay evidencia?"],
         ["Q08", "-1,95 (el más bajo de todas)", "sí, bien recuperada"],
         ["Q21", "-0,68", "no (control)"],
         ["Q19", "1,53", "sí"],
         ["Q13", "2,02", "sí"],
         ["Q22", "2,30", "no (control)"]]
tabla(s, datos, 0.7, 1.5, [2.2, 6.3, 3.4], alto=0.5, size=16)
bullets(s, [
    "Idea: no responder si el primer pasaje tiene un puntaje bajo",
    "No sirve: Q08 tiene evidencia y el puntaje más bajo, y Q22 queda por encima de Q19 y Q13",
    "  Un umbral que atrape Q21 también dejaría sin respuesta a Q08, que sí se podía responder",
    "No se implementó ningún umbral",
], y=4.7, h=2.4, size=19)

# 20. limitaciones
s = slide_titulo("Limitaciones")
bullets(s, [
    "12 consultas: sirve para ver casos concretos, no para generalizar",
    "Un solo modelo pequeño (Qwen2.5-1.5B) en CPU; uno más grande podría abstenerse mejor, y no se probó",
    "La auditoría la hicimos nosotros, sin un segundo anotador ni medida de acuerdo",
    "El corpus no tiene las bases del proceso, que están en PDF: consultas como Q22 no se pueden responder",
    "Una sola corrida y una sola máquina, así que las latencias son orientativas",
])

# 21. conclusion
s = slide_titulo("Conclusión")
bullets(s, [
    "Con evidencia, el mismo modelo pasa de no responder 9 de 10 consultas a responder bien 8 y a medias 1",
    "  Y 16 de sus 19 afirmaciones están respaldadas por el contexto",
    "Pero RAG le quitó la capacidad de abstenerse: 0 de 2 controles, frente a 2 de 2 en closed-book",
    "  El modelo usa los pasajes que recibe para inventar una respuesta con cita",
    "Una cita a un id válido no garantiza nada: 13 de 25 citas no sustentan lo que dice la respuesta",
    "  Una respuesta citada no es una respuesta verificada",
    "El costo: 3,5 veces más tokens de entrada y 75 % más de latencia",
], size=20)

salida = aqui / "Experimento2_EP.pptx"
prs.save(salida)
exportar(salida, aqui / "Experimento2_EP.pdf", aqui / "_png_exp2")
print("listo", salida, "| diapositivas:", len(prs.slides))
